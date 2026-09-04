#!/usr/bin/env python3
"""Build Hindi/Hinglish Depression Meme Dataset .docx from verified Pinterest posts."""

import hashlib
import json
import re
import sys
import time
import urllib.request
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent / ".vendor_py"))

from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image as PILImage
import io

WORK_DIR = Path(__file__).resolve().parent / "meme_dataset_build"
PINS_JSON = WORK_DIR / "verified_pins.json"
OUT_DOCX = Path(__file__).resolve().parent / "Hindi and Hinglish Depression Meme Dataset.docx"
IMG_DIR = WORK_DIR / "images"

EXCLUDE_URLS = {
    "https://in.pinterest.com/pin/742038476109636463/",  # shayari collection, not meme post
    "https://in.pinterest.com/pin/801640802441914933/",  # English promo / not Hindi meme
    "https://in.pinterest.com/pin/836051118356353036/",  # quote card, not meme
    "https://in.pinterest.com/pin/659636676672532195/",  # duplicate template of 824158800598945529
    "https://in.pinterest.com/pin/824158800598945529/",  # duplicate image of ye dukh template
}

CATEGORY_ORDER = [
    "Feeling Sad, Empty, or Hopeless",
    "Loss of Interest or Pleasure",
    "Sleep Problems",
    "Low Energy or Fatigue",
    "Appetite or Weight Changes",
    "Feelings of Worthlessness or Guilt",
    "Difficulty Concentrating",
    "Slowed Movement or Restlessness",
    "Thoughts About Death or Self-Harm",
]


def fetch(url, timeout=25):
    req = urllib.request.Request(
        url, headers={"User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36"}
    )
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return resp.read()


def pinterest_image_url(pin_url):
    html = fetch(pin_url).decode("utf-8", "ignore")
    m = re.search(r'property="og:image" content="([^"]+)"', html)
    if m:
        return m.group(1).replace("&amp;", "&")
    m = re.search(r"(https://i\.pinimg\.com/736x/[a-z0-9/]+\.(?:jpg|jpeg|png|gif|webp))", html)
    if m:
        return m.group(1)
    m = re.search(r"(https://i\.pinimg\.com/originals/[a-z0-9/]+\.(?:jpg|jpeg|png|gif|webp))", html)
    return m.group(1) if m else None


def categorize(title):
    t = title.lower()
    rules = [
        ("Thoughts About Death or Self-Harm", r"mar jau|kya karu|mar ja|suicide|self.?harm|mar jau"),
        ("Sleep Problems", r"neend|nind|insomnia|sleep cycle|raat|3 am|career ka soch"),
        ("Low Energy or Fatigue", r"thak|tired|fatigue|safai|hansa|khichdi|maine to|mai to thak"),
        ("Feelings of Worthlessness or Guilt", r"galti|bura|sorry sorry|itna galat|dictionary|worthless|guilty|narajgi"),
        ("Difficulty Concentrating", r"padhai|study|exam|school|cbse|kal karnege|begin|likhai|motivate"),
        ("Slowed Movement or Restlessness", r"raha nahi|dhak dhak|restless|anxiety|overthink|bol hi nahi|kya bolu"),
        ("Appetite or Weight Changes", r"bhook|appetite|weight|chawal|khana"),
        ("Loss of Interest or Pleasure", r"bore|barbaad hona|life se|struggle|interest|anhedonia|barbaad"),
        ("Feeling Sad, Empty, or Hopeless", r"dukh|udas|nirasha|hopeless|lonely|sapna toota|zindagi|sad|dard|jazbat|empty|depression|dispression|baigan|nahi meme|masaan|ye dukh|sala ye|overthinking|deep|shayari|2 asu|sticker|emoji|sabit|tasveer|mirzapur|munna bhai|jethalal"),
    ]
    for cat, pat in rules:
        if re.search(pat, t):
            return cat
    return "Feeling Sad, Empty, or Hopeless"


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(underline)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def verify_link(url):
    req = urllib.request.Request(
        url, method="HEAD", headers={"User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36"}
    )
    try:
        with urllib.request.urlopen(req, timeout=20) as resp:
            return resp.status in (200, 301, 302, 303, 307, 308)
    except Exception:
        try:
            fetch(url)
            return True
        except Exception:
            return False


def main():
    pins = json.loads(PINS_JSON.read_text())
    IMG_DIR.mkdir(parents=True, exist_ok=True)

    memes = []
    seen_hashes = set()

    for entry in pins:
        url = entry["url"]
        if url in EXCLUDE_URLS:
            continue
        if not verify_link(url):
            print(f"SKIP inaccessible: {url}")
            continue

        img_url = pinterest_image_url(url)
        if not img_url:
            print(f"SKIP no image: {url}")
            continue

        try:
            img_data = fetch(img_url)
        except Exception as e:
            print(f"SKIP download fail: {url} ({e})")
            continue

        img_hash = hashlib.sha256(img_data).hexdigest()
        if img_hash in seen_hashes:
            continue
        seen_hashes.add(img_hash)

        ext = Path(urllib.parse.urlparse(img_url).path).suffix.lower() or ".jpg"
        img_path = IMG_DIR / f"{img_hash[:16]}.png"
        try:
            with PILImage.open(io.BytesIO(img_data)) as im:
                im.convert("RGB").save(img_path, format="PNG")
        except Exception:
            if ext in {".jpg", ".jpeg", ".png"}:
                img_path = IMG_DIR / f"{img_hash[:16]}{ext}"
                img_path.write_bytes(img_data)
            else:
                print(f"SKIP unsupported image: {url}")
                continue

        memes.append(
            {
                "source": url,
                "image_path": img_path,
                "category": categorize(entry["title"]),
            }
        )
        print(f"OK {len(memes)}: {entry['title'][:60]}")
        time.sleep(0.1)

    print(f"Total verified memes: {len(memes)}")

    by_cat = {c: [] for c in CATEGORY_ORDER}
    for m in memes:
        by_cat[m["category"]].append(m)

    doc = Document()
    doc.add_heading("Hindi and Hinglish Depression Meme Dataset", 0)

    counter = 1
    for category in CATEGORY_ORDER:
        items = by_cat.get(category, [])
        if not items:
            continue
        doc.add_heading(category, level=1)
        for item in items:
            doc.add_paragraph(f"Meme {counter}")
            doc.add_picture(str(item["image_path"]), width=Inches(4.5))
            p = doc.add_paragraph("Source: ")
            add_hyperlink(p, item["source"], item["source"])
            doc.add_paragraph("")
            counter += 1

    doc.save(OUT_DOCX)
    print(f"Saved: {OUT_DOCX} ({counter - 1} memes)")


if __name__ == "__main__":
    import urllib.parse

    main()
