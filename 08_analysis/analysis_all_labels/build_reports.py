#!/usr/bin/env python3
"""Generate all-labels deep analysis HTML and Word doc."""

import base64
import html
import json
import re
import sys
from collections import Counter
from io import BytesIO
from pathlib import Path

from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parent))

ROOT = Path(__file__).resolve().parent.parent.parent
VENDOR = ROOT / "archive" / ".vendor_py"
if VENDOR.is_dir() and str(VENDOR) not in sys.path:
    sys.path.insert(0, str(VENDOR))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from label_config import LABELS7, MODEL_LABELS, SHORT

DATA_PATH = Path(__file__).resolve().parent / "data.json"
OUT_DIR = Path(__file__).resolve().parent


def replace_em_dash(text):
    if not text:
        return text
    text = re.sub(r"\s—\s", ", ", str(text))
    return text.replace("—", ", ")


def esc(s):
    return html.escape(replace_em_dash(str(s or "")))


def highlight_label(text, label):
    if not text:
        return '<em class="muted">(missing)</em>'
    text = replace_em_dash(text)
    from label_config import LABEL_PATTERNS

    patterns = LABEL_PATTERNS.get(label, {}).get("eng", []) + LABEL_PATTERNS.get(label, {}).get("deva", [])
    out = esc(text[:1000])
    for p in patterns:
        out = re.sub(f"({p})", r"<mark>\1</mark>", out, flags=re.I)
    return out.replace("\n", "<br>")


def load_image_bytes(rel_path, max_width=520):
    if not rel_path:
        return None
    p = ROOT / rel_path
    if not p.exists():
        return None
    try:
        raw = p.read_bytes()
    except (OSError, TimeoutError):
        return None
    try:
        im = Image.open(BytesIO(raw)).convert("RGB")
        if im.width > max_width:
            ratio = max_width / im.width
            im = im.resize((max_width, int(im.height * ratio)))
        buf = BytesIO()
        im.save(buf, format="JPEG", quality=80, optimize=True)
        return buf
    except Exception:
        return None


def img_data_uri(rel_path, max_width=480):
    stream = load_image_bytes(rel_path, max_width=max_width)
    if stream is None and rel_path:
        p = ROOT / rel_path
        if p.exists():
            try:
                raw = p.read_bytes()
                b64 = base64.b64encode(raw).decode("ascii")
                return f"data:image/jpeg;base64,{b64}"
            except (OSError, TimeoutError):
                return None
        return None
    if stream is None:
        return None
    b64 = base64.b64encode(stream.getvalue()).decode("ascii")
    return f"data:image/jpeg;base64,{b64}"


def english_image_path(sample_id):
    for rel in (
        f"dataset/test/{sample_id}.jpg",
        f"colab_experiments/data/test_images/{sample_id}.jpg",
    ):
        if (ROOT / rel).exists():
            return rel
    return None


def image_pair_html(ex):
    eng_rel = english_image_path(ex["sample_id"])
    deva_rel = ex.get("image_rel")
    eng_uri = img_data_uri(eng_rel)
    deva_uri = img_data_uri(deva_rel)
    eng_img = (
        f"<img src='{eng_uri}' alt='{esc(ex['sample_id'])} English'>"
        if eng_uri
        else "<p class='muted'>English image not found</p>"
    )
    deva_img = (
        f"<img src='{deva_uri}' alt='{esc(ex['sample_id'])} Devanagari'>"
        if deva_uri
        else "<p class='muted'>Devanagari image not found</p>"
    )
    return (
        "<div class='img-pair'>"
        f"<figure><figcaption class='eng'>Original (English)</figcaption>{eng_img}</figure>"
        f"<figure><figcaption class='deva'>Translated (Devanagari)</figcaption>{deva_img}</figure>"
        "</div>"
    )


def cond_group(key):
    c = key.split("/", 1)[1]
    if "image_only" in c:
        return "image_only"
    if "expl_only" in c:
        return "expl_only"
    return "multimodal"


def fmt_run(key):
    model, cond = key.split("/", 1)
    cond_n = cond.replace("_", " ").replace("deva", "Deva").replace("eng", "Eng").title()
    return f"{MODEL_LABELS.get(model, model)} ({cond_n})"


def wrong_labels(pred, target):
    return [l for l in LABELS7 if l != target and pred.get(l, 0) == 1]


def build_analysis(ex, label):
    short = SHORT[label]
    preds = ex.get("predictions", {})
    label_map = ex.get("label_predictions", {})
    n = len(label_map) or 1
    hits = [k for k, v in label_map.items() if v == 1]
    misses = [k for k, v in label_map.items() if v == 0]

    by_group = {"image_only": [], "expl_only": [], "multimodal": []}
    for k, v in label_map.items():
        by_group[cond_group(k)].append((k, v))

    wrong_counter = Counter()
    for k in misses:
        if k in preds:
            for l in wrong_labels(preds[k]["pred"], label):
                wrong_counter[l] += 1

    eng, deva = ex["eng_class"], ex["deva_class"]
    gold = ex["gold_labels"]
    pure = gold == [label]

    sections = []
    expl_points = []
    if eng == "explicit":
        expl_points.append(f"English explanation names {label}.")
    elif eng == "substitute":
        expl_points.append(f"English explanation substitutes {label} with a related symptom.")
    elif eng == "missing":
        expl_points.append("English explanation is missing or too short.")
    else:
        expl_points.append(f"English explanation does not signal {label}.")

    if deva == "explicit":
        expl_points.append(f"Hindi explanation names {label}.")
    elif deva == "substitute":
        expl_points.append(f"Hindi explanation emphasizes a related symptom, not {label}.")
    elif deva == "missing":
        expl_points.append("Hindi explanation is missing or too short.")
    else:
        expl_points.append(f"Hindi explanation omits {label} cues.")

    sections.append(
        "<h4>Explanation coverage</h4><ul>"
        + "".join(f"<li>{p}</li>" for p in expl_points)
        + "</ul>"
    )

    fail_pct = 100 * len(misses) / n
    hit_runs = ", ".join(fmt_run(k) for k in hits[:8]) if hits else "none"
    if len(hits) > 8:
        hit_runs += f", …and {len(hits) - 8} more"
    sections.append(
        f"<h4>Where models fail on {esc(label)}</h4>"
        f"<p><strong>{len(misses)}/{n} runs predict {short}=0</strong> "
        f"({fail_pct:.0f}% failure on this meme). "
        f"Runs that <span class='ok-text'>correctly predict {short}=1</span>: {esc(hit_runs)}.</p>"
    )

    group_lines = []
    for g, glabel in [
        ("image_only", "Image-only"),
        ("expl_only", "Explanation-only"),
        ("multimodal", "Image + explanation"),
    ]:
        items = by_group[g]
        if not items:
            continue
        ok = sum(1 for _, v in items if v == 1)
        group_lines.append(f"<li><strong>{glabel}:</strong> {ok}/{len(items)} correct</li>")
    sections.append("<ul class='compact'>" + "".join(group_lines) + "</ul>")

    if misses and wrong_counter:
        wrong_str = ", ".join(
            f"{l} ({c}/{len(misses)} misses)" for l, c in wrong_counter.most_common(3)
        )
        sections.append(
            f"<p>On failed runs, models most often fire <strong>wrong labels</strong> instead: {esc(wrong_str)}.</p>"
        )

    if misses:
        detail = []
        for k in sorted(misses)[:8]:
            wrong = wrong_labels(preds[k]["pred"], label) if k in preds else []
            wrong_s = ", ".join(wrong) if wrong else "no positive label"
            detail.append(
                f"<li>{esc(fmt_run(k))}: predicted <span class='bad-text'>{esc(wrong_s)}</span></li>"
            )
        if len(misses) > 8:
            detail.append(f"<li class='muted'>…and {len(misses) - 8} more failed runs</li>")
        sections.append(
            "<p><strong>Failed run details:</strong></p><ul class='compact'>"
            + "".join(detail)
            + "</ul>"
        )

    reasons = []
    if pure:
        reasons.append(
            f"Gold is {short}-only with no co-labels, so models cannot borrow cues from adjacent symptoms."
        )
    if eng != "explicit" and deva != "explicit":
        reasons.append(
            "Explanation pipeline failure: neither language surfaces the target symptom, "
            "so explanation-only and multimodal runs inherit the wrong clinical framing."
        )
    elif eng == "explicit" and len(misses) > len(hits):
        reasons.append(
            f"Explanation-classifier gap: English text mentions {label} but many runs still output {short}=0."
        )
    elif deva == "explicit" and any("deva_expl" in k for k in misses):
        reasons.append(
            f"Hindi explanation names {label} but some Devanagari runs still miss, "
            "possibly due to transliterated-label noise or co-occurring symptom confusion."
        )

    if by_group["image_only"] and all(v == 0 for _, v in by_group["image_only"]):
        reasons.append(
            f"Image-only failure: {label} is not visually or OCR-explicit in the Devanagari meme."
        )

    co_labels = [l for l in gold if l != label]
    if co_labels:
        reasons.append(
            f"Co-labeling with {', '.join(co_labels)}: salient cues for the co-label may satisfy models "
            f"before the subtler {label} dimension is tagged."
        )

    for wrong_label, count in wrong_counter.most_common(2):
        if count >= max(2, len(misses) // 3):
            reasons.append(
                f"Models often substitute {wrong_label} ({count}/{len(misses)} misses), "
                f"which overlaps semantically with {label} in meme text or explanations."
            )

    if ex.get("fail_rate", 1) <= 0.35:
        reasons.append(
            f"Most runs succeed when explanations or images make {label} explicit; "
            "remaining misses cluster on image-only or substitute-symptom inputs."
        )

    if not reasons:
        reasons.append(
            f"Mixed failure drivers: explanation substitution, image-only limits, and multi-label overlap."
        )

    sections.append(
        "<h4>Possible reasons for failure</h4><ul>"
        + "".join(f"<li>{r}</li>" for r in reasons)
        + "</ul>"
    )
    return "\n".join(sections)


CSS = """
body{font-family:Segoe UI,system-ui,sans-serif;background:#f8f9fb;color:#1e293b;line-height:1.55;margin:0;padding:24px;max-width:1100px;margin-inline:auto}
h1{color:#5b21b6} h2{border-bottom:2px solid #e2e8f0;padding-bottom:6px;margin-top:36px}
h3{margin-top:0} h4{margin:14px 0 6px;font-size:14px;color:#334155}
.label-section{margin-top:48px;padding-top:8px;border-top:3px solid #c4b5fd}
.card{background:#fff;border:1px solid #e2e8f0;border-radius:10px;padding:18px;margin:14px 0}
table{width:100%;border-collapse:collapse;font-size:14px} th,td{border:1px solid #e2e8f0;padding:6px 8px;text-align:center}
th{background:#f1f5f9} td:first-child,th:first-child{text-align:left}
.low{background:#fef2f2;color:#b91c1c;font-weight:600}.mid{background:#fffbeb;color:#b45309}
mark{background:#fef08a}.muted{color:#94a3b8;font-size:13px}
.badge{display:inline-block;padding:2px 8px;border-radius:999px;font-size:11px;font-weight:600;margin:2px}
.badge-gold{background:#ede9fe;color:#5b21b6}.badge-bad{background:#fee2e2;color:#b91c1c}.badge-ok{background:#dcfce7;color:#15803d}
.bad-text{color:#b91c1c;font-weight:600}.ok-text{color:#15803d;font-weight:600}
.case-grid{display:grid;grid-template-columns:420px 1fr;gap:16px;align-items:start}
.img-pair{display:grid;grid-template-columns:1fr 1fr;gap:10px;margin-bottom:10px}
.img-pair figure{margin:0}
.img-pair figcaption{font-size:11px;font-weight:600;color:#475569;text-align:center;margin-bottom:6px}
.img-pair img{width:100%;border-radius:8px;border:1px solid #e2e8f0;display:block}
.img-pair figcaption.deva{color:#c2410c}.img-pair figcaption.eng{color:#1d4ed8}
.expl-box{background:#f8fafc;border-left:4px solid #94a3b8;padding:10px 12px;margin:10px 0;font-size:13px;max-height:220px;overflow:auto}
.expl-box.eng{border-left-color:#2563eb}.expl-box.deva{border-left-color:#ea580c}
.analysis-box{background:#faf5ff;border:1px solid #e9d5ff;border-radius:8px;padding:12px 14px;margin-top:10px;font-size:13px}
.analysis-box ul{margin:6px 0;padding-left:18px}.analysis-box li{margin-bottom:5px}
ul.compact li{margin-bottom:3px}
.stat-grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:12px}
.stat{text-align:center;padding:14px;background:#f1f5f9;border-radius:8px}.stat .num{font-size:24px;font-weight:700;color:#5b21b6}
.finding li{margin-bottom:6px}
.pred-grid{display:flex;flex-wrap:wrap;gap:4px;margin:8px 0}
@media(max-width:800px){.case-grid{grid-template-columns:1fr}}
"""


def f1_table_html(f1_table):
    model_names = ["GPT-4o-mini", "Mistral-L3", "Mistral-S", "Llama-4", "Gemini"]
    parts = ["<table><tr><th>Condition</th>"]
    for m in model_names:
        parts.append(f"<th>{esc(m)}</th>")
    parts.append("</tr>")
    for row in f1_table:
        parts.append(f"<tr><td>{esc(row['condition'])}</td>")
        for m in model_names:
            v = row["scores"].get(m)
            if v is None:
                parts.append("<td>-</td>")
            else:
                cls = "low" if v < 0.4 else ("mid" if v < 0.55 else "")
                parts.append(f"<td class='{cls}'>{v:.4f}</td>")
        parts.append("</tr>")
    parts.append("</table>")
    return "".join(parts)


def stats_html(stats, label, short):
    gc = stats["gold_count"]
    eng = stats["eng"]
    deva = stats["deva"]
    explicit_eng = eng.get("explicit", 0)
    sub_eng = eng.get("substitute", 0)
    explicit_deva = deva.get("explicit", 0)
    both_miss = stats["both_miss"]
    return (
        "<div class='stat-grid'>"
        f"<div class='stat'><div class='num'>{explicit_eng}/{gc}</div>English mention {short}</div>"
        f"<div class='stat'><div class='num'>{sub_eng}/{gc}</div>English substitute</div>"
        f"<div class='stat'><div class='num'>{explicit_deva}/{gc}</div>Hindi mention {short}</div>"
        f"<div class='stat'><div class='num'>{both_miss}/{gc}</div>Both miss/substitute</div>"
        "</div>"
    )


def build_html(data):
    parts = [
        "<!DOCTYPE html><html lang='en'><head><meta charset='UTF-8'>",
        "<meta name='viewport' content='width=device-width, initial-scale=1'>",
        "<title>All Labels Deep Analysis</title><style>", CSS, "</style></head><body>",
        "<h1>PHQ-9 Symptom Labels: Deep Analysis (All 7 Labels)</h1>",
        "<p style='color:#64748b'>Self-contained report with embedded images. "
        f"Based on {data['run_count']} model runs across {data['sample_count']} test memes. "
        "Each label section includes F1 breakdown, explanation coverage, and 3 case studies with "
        "English/Devanagari images, explanations, per-run predictions, and failure analysis.</p>",
    ]

    for idx, lb in enumerate(data["labels"], 1):
        label = lb["name"]
        short = lb["short"]
        parts.append(f"<div class='label-section' id='label-{idx}'>")
        parts.append(f"<h2>{idx}. {esc(label)} ({short})</h2>")

        parts.append("<h3>Per-condition F1</h3><div class='card'>")
        parts.append(f1_table_html(lb["f1_table"]))
        parts.append(f"<p>{esc(lb['global_summary'])}</p></div>")

        parts.append(
            f"<h3>Explanation coverage ({lb['gold_count']} gold {short} memes)</h3>"
            f"<div class='card'>{stats_html(lb['stats'], label, short)}</div>"
        )

        parts.append("<h3>Case studies</h3>")
        for ex in lb["examples"]:
            gold_badges = "".join(
                f"<span class='badge badge-gold'>{esc(l)}</span>" for l in ex["gold_labels"]
            )
            pred_bits = []
            for key, v in sorted(ex.get("label_predictions", {}).items()):
                cls = "badge-ok" if v == 1 else "badge-bad"
                pred_bits.append(
                    f"<span class='badge {cls}'>{esc(key.replace('/', ' '))}: {short}={v}</span>"
                )

            parts.append(f"<div class='card'><h3>{esc(ex['title'])}</h3>")
            parts.append(f"<p>{esc(ex['category'])} | Gold: {gold_badges}</p>")
            parts.append("<div class='case-grid'>")
            ocr = ex.get("ocr_text", "")[:180]
            parts.append(
                f"<div>{image_pair_html(ex)}"
                f"<p class='muted'><b>Original OCR:</b> {esc(ocr)}...</p></div>"
            )
            parts.append("<div>")
            parts.append("<p><b>English explanation</b> <span class='muted'>(mistral_large3/eng)</span></p>")
            parts.append(f"<div class='expl-box eng'>{highlight_label(ex['eng_explanation'], label)}</div>")
            parts.append("<p><b>Hindi explanation</b> <span class='muted'>(gpt4o_mini/deva)</span></p>")
            parts.append(f"<div class='expl-box deva'>{highlight_label(ex['deva_explanation'], label)}</div>")
            parts.append(f"<p><b>{short} predictions ({ex['hits']}/{ex['total_runs']} hits):</b></p>")
            parts.append(f"<div class='pred-grid'>{''.join(pred_bits)}</div>")
            parts.append(f"<div class='analysis-box'>{build_analysis(ex, label)}</div>")
            parts.append("</div></div></div>")

        parts.append("<h3>Recommendations</h3><div class='card'><ul class='finding'>")
        for rec in lb["recommendations"]:
            parts.append(f"<li>{esc(rec)}</li>")
        parts.append("</ul></div></div>")

    parts.append("</body></html>")
    return "".join(parts)


def add_heading(doc, text, level=1):
    return doc.add_heading(replace_em_dash(text), level=level)


def add_para(doc, text, bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(replace_em_dash(text))
    run.font.size = Pt(size)
    run.bold = bold
    return p


def add_images(doc, ex):
    eng_rel = english_image_path(ex["sample_id"])
    deva_rel = ex.get("image_rel")
    for caption, rel in [("Original (English)", eng_rel), ("Translated (Devanagari)", deva_rel)]:
        stream = load_image_bytes(rel)
        if stream is None:
            if rel:
                add_para(doc, f"{caption}: (image unavailable)")
            continue
        try:
            stream.seek(0)
            doc.add_paragraph().add_run(f"{caption}:").bold = True
            doc.add_picture(stream, width=Inches(2.2))
        except Exception:
            add_para(doc, f"{caption}: (image unavailable)")


def f1_table_docx(doc, f1_table):
    model_names = ["GPT-4o-mini", "Mistral-L3", "Mistral-S", "Llama-4", "Gemini"]
    table = doc.add_table(rows=1 + len(f1_table), cols=1 + len(model_names))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Condition"
    for i, m in enumerate(model_names, 1):
        hdr[i].text = m
    for ri, row in enumerate(f1_table, 1):
        cells = table.rows[ri].cells
        cells[0].text = row["condition"]
        for ci, m in enumerate(model_names, 1):
            v = row["scores"].get(m)
            cells[ci].text = f"{v:.4f}" if v is not None else "-"


def build_docx(data):
    doc = Document()
    title = doc.add_heading("PHQ-9 Symptom Labels: Deep Analysis (All 7 Labels)", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_para(
        doc,
        f"Based on {data['run_count']} model runs across {data['sample_count']} test memes. "
        "Each label includes F1 breakdown, explanation coverage, 3 case studies, and recommendations.",
    )

    for idx, lb in enumerate(data["labels"], 1):
        label = lb["name"]
        short = lb["short"]
        doc.add_page_break()
        add_heading(doc, f"{idx}. {label} ({short})", 1)

        add_heading(doc, "Per-condition F1", 2)
        f1_table_docx(doc, lb["f1_table"])
        add_para(doc, lb["global_summary"])

        add_heading(doc, f"Explanation coverage ({lb['gold_count']} gold {short} memes)", 2)
        st = lb["stats"]
        gc = st["gold_count"]
        add_para(
            doc,
            f"English explicit: {st['eng'].get('explicit', 0)}/{gc} | "
            f"English substitute: {st['eng'].get('substitute', 0)}/{gc} | "
            f"Hindi explicit: {st['deva'].get('explicit', 0)}/{gc} | "
            f"Both miss/substitute: {st['both_miss']}/{gc}",
        )

        add_heading(doc, "Case studies", 2)
        for ex in lb["examples"]:
            add_heading(doc, ex["title"], 3)
            gold = ", ".join(ex["gold_labels"])
            add_para(doc, f"{ex['category']} | Gold: {gold}")
            add_images(doc, ex)
            ocr = ex.get("ocr_text", "")[:200]
            add_para(doc, f"Original OCR: {ocr}...", bold=True)
            add_para(doc, "English explanation (mistral_large3/eng):", bold=True)
            add_para(doc, ex.get("eng_explanation", "")[:1200])
            add_para(doc, "Hindi explanation (gpt4o_mini/deva):", bold=True)
            add_para(doc, ex.get("deva_explanation", "")[:1200])

            add_para(doc, f"{short} predictions ({ex['hits']}/{ex['total_runs']} hits):", bold=True)
            for key, v in sorted(ex.get("label_predictions", {}).items()):
                add_para(doc, f"  {key.replace('/', ' ')}: {short}={v}")

            analysis_html = build_analysis(ex, label)
            add_para(doc, "Analysis:", bold=True)
            for block in re.findall(r"<li>(.*?)</li>", analysis_html):
                add_para(doc, f"• {html.unescape(re.sub('<[^>]+>', '', block))}")

        add_heading(doc, "Recommendations", 2)
        for rec in lb["recommendations"]:
            add_para(doc, f"• {rec}")

    return doc


def main():
    import shutil
    import sys

    docx_only = "--docx-only" in sys.argv
    data = json.loads(DATA_PATH.read_text(encoding="utf-8"))

    if not docx_only:
        html_content = build_html(data)
        html_path = OUT_DIR / "all_labels_deep_analysis.html"
        html_path.write_text(html_content, encoding="utf-8")
        print(f"Wrote {html_path} ({html_path.stat().st_size // 1024} KB)")

        standalone = OUT_DIR / "all_labels_deep_analysis_standalone.html"
        standalone.write_text(html_content, encoding="utf-8")

        desktop_html = Path.home() / "Desktop" / "All_Labels_Deep_Analysis.html"
        desktop_html.write_text(html_content, encoding="utf-8")

    doc = build_docx(data)
    docx_path = OUT_DIR / "All_Labels_Deep_Analysis.docx"
    doc.save(docx_path)
    print(f"Wrote {docx_path} ({docx_path.stat().st_size // 1024} KB)")

    desktop_docx = Path.home() / "Desktop" / "All_Labels_Deep_Analysis.docx"
    shutil.copy2(docx_path, desktop_docx)
    print("Copied to Desktop")


if __name__ == "__main__":
    main()
